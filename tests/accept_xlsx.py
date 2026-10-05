"""Acceptance test for the calculator's .xlsx output, using openpyxl.

Verifies structure, numeric typing, the odometer chain and the total, and
exits non-zero on any failure so it can gate a release.
"""
import os
import re
import sys
import zipfile

from lxml import etree
from openpyxl import load_workbook

path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "out", "Travel_Log.xlsx")
ONEWAY = round(float(sys.argv[2]), 2) if len(sys.argv) > 2 else 45.0
DAYS = int(sys.argv[3]) if len(sys.argv) > 3 else 5
START_ODO = float(sys.argv[4]) if len(sys.argv) > 4 else 92451.0

passed = failed = 0
TOL = 0.011   # values are stored to two decimal places


def check(name, cond, detail=""):
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}" + (f" -> {detail}" if detail else ""))


def eq(name, actual, expected):
    check(name, actual == expected, f"got {actual!r}, expected {expected!r}")


print(f"\n== {path} (one-way {ONEWAY}, {DAYS} days, start odo {START_ODO}) ==")

# --- package integrity -----------------------------------------------------
with zipfile.ZipFile(path) as z:
    bad = z.testzip()
    check("zip integrity (all CRCs)", bad is None, f"first bad entry: {bad}")
    for name in z.namelist():
        if name.endswith((".xml", ".rels")):
            try:
                etree.fromstring(z.read(name))
                check(f"well-formed XML: {name}", True)
            except Exception as exc:
                check(f"well-formed XML: {name}", False, str(exc))

wb = load_workbook(path)
eq("sheet name", wb.sheetnames, ["Travel Log"])
ws = wb.active

# --- locate the table ------------------------------------------------------
HEADER = None
for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    if row[0].value == "DATE" and row[8].value == "KILOMETRES":
        HEADER = row[0].row
        break
check("found the table header row", HEADER is not None)
if HEADER is None:
    sys.exit(1)
print(f"  info  table header on row {HEADER}")

trip_rows = []
r = HEADER + 1
while r <= ws.max_row and isinstance(ws[f"A{r}"].value, str) and re.match(r"^\d{2}-[A-Z][a-z]{2}-\d{4}$", ws[f"A{r}"].value):
    trip_rows.append(r)
    r += 1

eq("trip row count = days x 2", len(trip_rows), DAYS * 2)
print(f"  info  trip rows {trip_rows[0]}..{trip_rows[-1]}")

# --- numeric typing --------------------------------------------------------
text_cols = []
for col in ("G", "H", "I"):
    for rr in trip_rows:
        v = ws[f"{col}{rr}"].value
        if v is None or isinstance(v, str):
            text_cols.append(f"{col}{rr}={v!r}")
check("all odometer/kilometre cells numeric", not text_cols, ", ".join(text_cols[:5]))

# --- odometer chain --------------------------------------------------------
expected_odo = START_ODO
chain_errors = []
for i, rr in enumerate(trip_rows):
    g, h, km = ws[f"G{rr}"].value, ws[f"H{rr}"].value, ws[f"I{rr}"].value
    if abs(g - expected_odo) > TOL:
        chain_errors.append(f"row {rr}: start {g} != expected {expected_odo}")
    if abs(h - (g + ONEWAY)) > TOL:
        chain_errors.append(f"row {rr}: end {h} != start {g} + {ONEWAY}")
    if abs(km - ONEWAY) > TOL:
        chain_errors.append(f"row {rr}: km {km} != {ONEWAY}")
    expected_odo = h
check("odometer chains correctly across all legs", not chain_errors, "; ".join(chain_errors[:4]))
check("final odometer",
      abs(ws[f"H{trip_rows[-1]}"].value - (START_ODO + DAYS * 2 * ONEWAY)) <= TOL,
      f"got {ws[f'H{trip_rows[-1]}'].value}, expected {START_ODO + DAYS * 2 * ONEWAY}")

# --- total -----------------------------------------------------------------
total_label_row = r
eq("total label row", ws[f"A{total_label_row}"].value, "TOTAL KILOMETRES TRAVELLED")
total_cell = ws[f"I{total_label_row}"].value
eq("total is a SUM over the trip rows",
   total_cell, f"=SUM(I{trip_rows[0]}:I{trip_rows[-1]})")
expected_total = DAYS * 2 * ONEWAY
_cell_sum = round(sum(ws[f"I{rr}"].value for rr in trip_rows), 2)
check("sum of km cells equals expected total",
      abs(_cell_sum - expected_total) <= TOL,
      f"got {_cell_sum}, expected {expected_total}")

# cached value inside the sheet XML (so the figure shows before recalc)
with zipfile.ZipFile(path) as z:
    sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
m = re.search(r'<c r="I%d"[^>]*><f>[^<]*</f><v>([^<]*)</v></c>' % total_label_row, sheet)
check("total carries a cached value", m is not None)
if m:
    check("cached total value", abs(float(m.group(1)) - expected_total) <= TOL,
          f"got {m.group(1)}, expected {expected_total}")

# --- content sanity --------------------------------------------------------
check("no cell contains a raw newline", "\\n" not in sheet and "\n" not in sheet.split("<sheetData>")[1])
check("dates present", ws[f"A{trip_rows[0]}"].value is not None)
check("table has 9 columns", ws.max_column == 9)
check("mergeCells declared", "<mergeCells" in sheet)
check("autofilter covers the data", f'autoFilter ref="A{HEADER}:I{trip_rows[-1]}"' in sheet)

print(f"\n================ {passed} passed, {failed} failed ================\n")
sys.exit(1 if failed else 0)
